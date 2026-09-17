from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from nfp_service import exportar_rateio_xlsx
from nfp_utils import formatar_cnpj


def test_formatar_cnpj_mascara_completa():
    assert formatar_cnpj("11222333000181") == "11.222.333/0001-81"
    assert formatar_cnpj("11.222.333/0001-81") == "11.222.333/0001-81"
    assert formatar_cnpj("123") == "123"
    assert formatar_cnpj(None) == ""


def test_exportar_rateio_xlsx_cnpj_como_texto():
    linhas = [
        SimpleNamespace(
            cnpj="11222333000181",
            loja="Loja Teste",
            captador="DIEGO",
            origem="DIEGO",
            qtd=2,
            retorno=10.5,
            valor_diego=5.25,
            valor_aeb=5.25,
            final=10.5,
            competencia="2026-09",
        )
    ]
    conteudo = exportar_rateio_xlsx(linhas)
    wb = load_workbook(BytesIO(conteudo))
    ws = wb.active
    assert ws["A2"].value == "11.222.333/0001-81"
    assert ws["A2"].number_format == "@"
    assert ws["B2"].value == "Loja Teste"
